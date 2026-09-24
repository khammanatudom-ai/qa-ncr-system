import os
import io
import openpyxl
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

def fill_customer1_excel(uploaded_excel_bytes, ai_summary_data):
    """หยอดข้อมูลลงในช่องสีเขียวของ Customer 1 (U-SHIN)"""
    wb = openpyxl.load_workbook(io.BytesIO(uploaded_excel_bytes))
    sheet_name = 'U-SHIN' if 'U-SHIN' in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet_name]
    
    # หยอดข้อมูลลงตำแหน่งเซลล์
    ws['AB87'] = ai_summary_data.get('summary_occ', '-')
    ws['AB127'] = ai_summary_data.get('summary_outflow', '-')
    ws['C86'] = ai_summary_data.get('why_occ', '-')
    ws['AA115'] = ai_summary_data.get('why_outflow', '-')
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def fill_customer3_excel(uploaded_excel_bytes, ai_summary_data):
    """หยอดข้อมูลลงใน 3 ช่องของ Customer 3"""
    wb = openpyxl.load_workbook(io.BytesIO(uploaded_excel_bytes))
    ws = wb.active
    
    ws['C26'] = ai_summary_data.get('cause_of_problem', '-')
    ws['C34'] = ai_summary_data.get('why_undetected', '-')
    ws['C42'] = ai_summary_data.get('corrective_action', '-')
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def fill_ncr_pptx_template(arg1="why why analysis form Rev.3 use this.pptx", arg2=None, image_file=None, template_path="why why analysis form Rev.3 use this.pptx", ncr_data=None, output_dir="NCR_Issued", ncr_no=None):
    if isinstance(arg1, dict): ncr_data = arg1
    elif isinstance(arg1, str) and (arg1.endswith(".pptx") or os.path.exists(arg1)): template_path = arg1

    if isinstance(arg2, dict): ncr_data = arg2
    if ncr_data is None: ncr_data = {}

    if not os.path.exists(template_path):
        if os.path.exists("why why analysis form Rev.3 use this.pptx"):
            template_path = "why why analysis form Rev.3 use this.pptx"

    prs = Presentation(template_path)

    doc_no = str(ncr_no or ncr_data.get("doc_no") or ncr_data.get("ncr_no") or "-")
    prob_date = str(ncr_data.get("occur_date") or ncr_data.get("problem_date") or "-")
    model_name = str(ncr_data.get("model_name") or ncr_data.get("model") or "-")
    part_no = str(ncr_data.get("part_no", "-"))
    part_name = str(ncr_data.get("part_name", "-"))
    supplier = str(ncr_data.get("supplier", "-"))
    problem = str(ncr_data.get("problem", "-"))
    detail = str(ncr_data.get("problem_details") or ncr_data.get("detail", "-"))

    replacements = {
        "{{DOC_NO}}": doc_no,
        "{{OCCUR_DATE}}": prob_date,
        "{{SUPPLIER}}": supplier,
        "{{PART_NO}}": part_no,
        "{{PART_NAME}}": part_name,
        "{{MODEL_NAME}}": model_name,
        "{{PROBLEM}}": problem,
        "{{PROBLEM_DETAILS}}": detail,
    }

    def replace_in_paragraphs(paragraphs):
        for paragraph in paragraphs:
            for key, val in replacements.items():
                if key in paragraph.text:
                    paragraph.text = paragraph.text.replace(key, val)

    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                replace_in_paragraphs(shape.text_frame.paragraphs)
            if shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        replace_in_paragraphs(cell.text_frame.paragraphs)

    pptx_out = io.BytesIO()
    prs.save(pptx_out)
    pptx_out.seek(0)
    return pptx_out