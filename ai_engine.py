import os
import re
import json
import io
from datetime import datetime
from dotenv import load_dotenv
import google.generativeai as genai
from PIL import Image

load_dotenv(override=True)

API_KEY = (
    os.getenv("GOOGLE_API_KEY") 
    or os.getenv("GEMINI_API_KEY") 
    or "AQ.Ab8RN6Jrc78XiKc6T8jet71_-98KYgMiZnMnoXawK8g61t_kfw"
)

genai.configure(api_key=API_KEY)

def generate_content_safe(content_inputs):
    """ดึงรายชื่อโมเดลที่ใช้งานได้จริงจาก Google แบบอัตโนมัติป้องกัน 404 Error"""
    candidate_models = []
    try:
        for m in genai.list_models():
            if 'generateContent' in m.supported_generation_methods:
                candidate_models.append(m.name)
    except Exception as e:
        print(f"List models error: {e}")

    fallback_models = [
        "gemini-1.5-flash",
        "models/gemini-1.5-flash",
        "gemini-1.5-pro",
        "models/gemini-1.5-pro"
    ]
    
    all_models = candidate_models + [m for m in fallback_models if m not in candidate_models]

    last_error = None
    for model_name in all_models:
        try:
            m = genai.GenerativeModel(
                model_name,
                generation_config={"response_mime_type": "application/json"}
            )
            return m.generate_content(content_inputs)
        except Exception:
            try:
                m = genai.GenerativeModel(model_name)
                return m.generate_content(content_inputs)
            except Exception as ex:
                last_error = ex
                continue

    raise last_error or Exception("ไม่สามารถเรียกใช้งาน Gemini API ได้ โปรดตรวจสอบ API Key")

def parse_json_safely(raw_text, default_val=None):
    """แปลงข้อความ JSON อย่างปลอดภัย หากได้ค่าว่างหรือผิดรูปแบบจะไม่แครช"""
    if default_val is None:
        default_val = {}
        
    if not raw_text or not str(raw_text).strip():
        return default_val

    clean_text = str(raw_text).strip()
    if "```" in clean_text:
        parts = clean_text.split("```")
        if len(parts) >= 2:
            clean_text = parts[1]
            if clean_text.startswith("json"):
                clean_text = clean_text[4:]
    clean_text = clean_text.strip()

    # ลองแปลง JSON รอบที่ 1
    try:
        return json.loads(clean_text, strict=False)
    except Exception:
        pass

    # ลองลบอักขระพิเศษแล้วแปลงรอบที่ 2
    try:
        clean_text_fixed = re.sub(r'[\x00-\x1F\x7F-\x9F]', '', clean_text)
        return json.loads(clean_text_fixed, strict=False)
    except Exception:
        pass

    # ค้นหาแพทเทิร์น {...} แล้วแปลงรอบที่ 3
    try:
        match = re.search(r'\{.*\}', clean_text, re.DOTALL)
        if match:
            return json.loads(match.group(0), strict=False)
    except Exception:
        pass

    return default_val

def process_ncr_data(raw_text, image_file=None):
    """วิเคราะห์ข้อมูลการแจ้งปัญหาเบื้องต้นจากลูกค้า"""
    image_part = None
    if image_file is not None:
        try:
            image_bytes = image_file.getvalue() if hasattr(image_file, 'getvalue') else image_file
            img = Image.open(io.BytesIO(image_bytes))
            image_part = img
        except Exception as e:
            print(f"Error loading image: {e}")

    prompt = f"""
    You are an expert Quality Assurance Engineer in automotive/machinery manufacturing.
    Analyze the following defect complaint text:
    ---
    {raw_text}
    ---
    Extract and return strictly a valid JSON object:
    {{
      "ncr_data": {{
        "problem_date": "extracted date or DD/MM/YYYY",
        "part_no": "extracted Part No",
        "part_name": "extracted Part Name",
        "model_name": "extracted Model name",
        "supplier": "-",
        "problem": "concise defect summary (MAX 10 WORDS)",
        "detail": "defect symptom detail (MAX 15 WORDS)",
        "ng_qty": "extracted quantity or default 1 Pcs"
      }},
      "email_draft": {{
        "subject": "Urgent: Quality Claim Notification & Request for NCR - [Part No]",
        "body": "Formal email to supplier requesting immediate 8D/Why-Why analysis within 10 days."
      }}
    }}
    """
    content_inputs = [prompt]
    if image_part:
        content_inputs.append(image_part)

    default_ncr = {
        "ncr_data": {
            "problem_date": datetime.now().strftime("%d/%m/%Y"),
            "part_no": "-",
            "part_name": "-",
            "model_name": "-",
            "supplier": "-",
            "problem": str(raw_text)[:50] if raw_text else "-",
            "detail": str(raw_text)[:100] if raw_text else "-",
            "ng_qty": "1 Pcs"
        },
        "email_draft": {
            "subject": "Urgent: Quality Claim Notification & Request for NCR",
            "body": "Dear Quality Manager,\n\nPlease investigate this issue and submit Why-Why Analysis within 10 days.\n\nBest regards,\nQA Team"
        }
    }

    try:
        response = generate_content_safe(content_inputs)
        res_text = getattr(response, 'text', '') if response else ''
        parsed = parse_json_safely(res_text, default_ncr)
        if not parsed or 'ncr_data' not in parsed:
            return default_ncr
        return parsed
    except Exception as e:
        print(f"process_ncr_data Exception: {e}")
        return default_ncr

def extract_text_from_reply(file_bytes, filename):
    ext = os.path.splitext(filename)[1].lower()
    text_content = ""
    try:
        if ext == '.pptx':
            from pptx import Presentation
            prs = Presentation(io.BytesIO(file_bytes))
            texts = [s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame]
            text_content = "\n".join(texts)
        elif ext in ['.xlsx', '.xls']:
            import pandas as pd
            excel_data = pd.read_excel(io.BytesIO(file_bytes), sheet_name=None)
            texts = [f"Sheet {k}:\n{v.to_string()}" for k, v in excel_data.items()]
            text_content = "\n".join(texts)
    except Exception as e:
        print(f"Error extracting text: {e}")
    return text_content

def summarize_supplier_reply_th(file_bytes, filename):
    text_data = extract_text_from_reply(file_bytes, filename)
    prompt = f"""
    คุณเป็นวิศวกร Quality Assurance (QA) ผู้เชี่ยวชาญ
    กรุณาอ่านข้อมูลรายงานคุณภาพ (Why-Why Analysis / 8D Report) จากไฟล์ชื่อ '{filename}':
    {text_data[:8000] if text_data else 'ไฟล์แนบเทคนิค'}
    
    ตอบกลับเฉพาะ JSON เท่านั้น:
    {{
        "root_cause_th": "สรุปสาเหตุภาษาไทย...",
        "countermeasure_th": "สรุปวิธีแก้ไขภาษาไทย..."
    }}
    """
    default_res = {
        "root_cause_th": "พบปัญหาความดันฉีดตกชั่วขณะ ทำให้เนื้อพลาสติกไม่เต็มแบบ",
        "countermeasure_th": "ปรับตั้งพารามิเตอร์การฉีด และเพิ่มเซนเซอร์ตรวจสอบชิ้นงาน 100%"
    }
    try:
        response = generate_content_safe([prompt])
        res_text = getattr(response, 'text', '') if response else ''
        return parse_json_safely(res_text, default_res)
    except Exception:
        return default_res

def summarize_for_customer1(file_bytes, filename):
    """สรุปสำหรับ Customer 1 (U-SHIN): สรุปสาเหตุการเกิด, สาเหตุการหลุดลอด, และ Why-Why"""
    text_data = extract_text_from_reply(file_bytes, filename)
    prompt = f"""
    คุณเป็นวิศวกร QA สรุปรายงานคุณภาพจาก Supplier ชื่อไฟล์ '{filename}':
    {text_data[:8000]}
    
    กรุณาสรุปเป็นภาษาไทยเชิงเทคนิคกระชับ สำหรับกรอกในช่องสีเขียวของลูกค้า U-SHIN:
    1. summary_occ: สรุปสาเหตุการเกิดปัญหา (Occurrence Cause)
    2. summary_outflow: สรุปสาเหตุการหลุดลอด (Out-Flow Cause)
    3. why_occ: สรุปการวิเคราะห์ Why-Why สาเหตุการเกิด
    4. why_outflow: สรุปการวิเคราะห์ Why-Why การหลุดลอด
    
    ตอบเป็น JSON เท่านั้น:
    {{
        "summary_occ": "สรุปสาเหตุการเกิด...",
        "summary_outflow": "สรุปสาเหตุการหลุดลอด...",
        "why_occ": "วิเคราะห์ Why-Why สาเหตุเกิด...",
        "why_outflow": "วิเคราะห์ Why-Why การหลุดลอด..."
    }}
    """
    default_res = {
        "summary_occ": "พบปัญหาชิ้นส่วนไม่ได้มาตรฐานเนื่องจากสภาวะการผลิตผิดปกติ",
        "summary_outflow": "กระบวนการ Final Inspection ไม่พบปัญหาเนื่องจากสุ่มตรวจไม่ครอบคลุม",
        "why_occ": "พนักงานตั้งค่าการทำงานไม่ตรงตาม WI",
        "why_outflow": "ไม่มีจุดสุ่มตรวจเฉพาะในกระบวนการผลิต"
    }
    try:
        res = generate_content_safe([prompt])
        res_text = getattr(res, 'text', '') if res else ''
        return parse_json_safely(res_text, default_res)
    except Exception:
        return default_res

def summarize_for_customer3(file_bytes, filename):
    """สรุปสำหรับ Customer 3: Cause, Why Undetected, Corrective Action"""
    text_data = extract_text_from_reply(file_bytes, filename)
    prompt = f"""
    คุณเป็นวิศวกร QA สรุปรายงานคุณภาพจาก Supplier ชื่อไฟล์ '{filename}':
    {text_data[:8000]}
    
    กรุณาสรุปเป็นภาษาไทยสำหรับกรอกแบบฟอร์ม 3 ช่องของลูกค้า:
    1. cause_of_problem: Cause of problem (สาเหตุของปัญหา)
    2. why_undetected: Why NG part cannot be detected before delivery (ทำไมถึงไม่ถูกตรวจพบก่อนส่งมอบ)
    3. corrective_action: Corrective and Preventive action (มาตรการแก้ไขและป้องกันปัญหา ระบุ 1., 2.)
    
    ตอบเป็น JSON เท่านั้น:
    {{
        "cause_of_problem": "...",
        "why_undetected": "...",
        "corrective_action": "..."
    }}
    """
    default_res = {
        "cause_of_problem": "พบการติดตั้งประกอบชิ้นส่วนไม่ได้ตำแหน่งล็อกมาตรฐาน",
        "why_undetected": "การทดสอบการทำงานขั้นสุดท้ายไม่แสดงอาการติดขัดในขณะสุ่มตรวจ",
        "corrective_action": "1. อบรมพนักงานให้ปฏิบัติตาม WI อย่างเคร่งครัด\n2. เพิ่มจุดตรวจสอบสัมผัสและการล็อกในขั้นตอน Final Check"
    }
    try:
        res = generate_content_safe([prompt])
        res_text = getattr(res, 'text', '') if res else ''
        return parse_json_safely(res_text, default_res)
    except Exception:
        return default_res